"""캉골 RT 검토맵 데이터 생성기.

RT 폴더의 최신 `캉골RT내역_YYMMDD.xlsx`(raw)를 읽어 `data/rt.js`를 만든다.
상품 정보(품명·기획·복종·최초가)와 이미지는 ../product-info 기초데이터에서 가져온다.

사용: python build_rt.py [raw파일경로] [--push]
  --push  생성 후 GitHub에 커밋·푸시 (RT 저장소 + product-info 이미지 변경분)
"""
import datetime as dt
import glob
import json
import os
import re
import shutil
import subprocess
import sys
from collections import defaultdict

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
PRODUCT_INFO = os.path.normpath(os.path.join(HERE, '..', 'product-info'))
OUT = os.path.join(HERE, 'data', 'rt.js')

ALL_TYPES = ('백화점', '아울렛', '복합쇼핑몰')
CO_TYPES = ('아울렛', '복합쇼핑몰')
SIZE_ORDER = ['XS', 'SS', 'ME', 'LA', 'XL', 'XX', '2X', '3X', 'OS']


def cat_of(plan):
    plan = plan or ''
    if '모자' in plan:
        return '모자'
    if '의류' in plan:
        return '의류'
    if '가방' in plan:
        return '가방'
    return None


def num(v):
    if v is None or v == '':
        return 0
    try:
        f = float(v)
    except (TypeError, ValueError):
        return 0
    return int(f) if f == int(f) else f


def to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if v is None:
        return None
    s = str(v).strip()[:10].replace('.', '-').replace('/', '-')
    try:
        return dt.date.fromisoformat(s)
    except ValueError:
        return None


def date_str(v):
    d = to_date(v)
    return d.isoformat() if d else None


def season_of(year, season_name):
    y = str(year or '').strip()[-2:]
    s = str(season_name or '')
    return y + ('AW' if ('가을' in s or '겨울' in s or 'AW' in s.upper()) else 'SS')


class Sheet:
    """헤더 행을 자동으로 찾아 컬럼명으로 값을 읽는다. 같은 이름이 여러 번 나오면 상위 구분행 이름으로 구별한다."""

    def __init__(self, ws, key='수입품번'):
        self.title = ws.title
        rows = list(ws.iter_rows(values_only=True))
        hi = next(i for i, r in enumerate(rows[:15]) if key in [str(c).strip() if c else c for c in r])
        self.header = [str(c).strip() if c is not None else '' for c in rows[hi]]
        group, last = [], ''
        above = rows[hi - 1] if hi > 0 else [None] * len(self.header)
        for i in range(len(self.header)):
            g = above[i] if i < len(above) else None
            if g not in (None, ''):
                last = str(g).strip()
            group.append(last)
        self.group = group
        self.rows = [r for r in rows[hi + 1:] if r and any(c not in (None, '') for c in r)]

    def col(self, name, group=None):
        for i, h in enumerate(self.header):
            if h == name and (group is None or self.group[i] == group):
                return i
        raise KeyError(f'[{self.title}] 컬럼 "{name}"{f" (구분 {group})" if group else ""} 없음')

    def cols(self, *names):
        return [self.col(n) for n in names]


def find_sheet(wb, pred, label):
    for ws in wb.worksheets:
        if pred(ws.title.strip()):
            return ws
    raise KeyError(f'시트를 찾을 수 없음: {label} (시트 목록: {[w.title for w in wb.worksheets]})')


def parse_events(sheet, year):
    """'10.2 ~ 10.8' 형태(연도 없음)를 판매기간 연도로 해석."""
    ci_store, ci_period = sheet.col('매장명'), sheet.col('기간')
    out = []
    for r in sheet.rows:
        store, period = r[ci_store], r[ci_period]
        if not store or not period:
            continue
        m = re.match(r'\s*(\d{1,2})\s*[./]\s*(\d{1,2})\s*~\s*(\d{1,2})\s*[./]\s*(\d{1,2})', str(period))
        if not m:
            print(f'  ! 행사기간 형식 인식 불가: {store} "{period}"')
            continue
        m1, d1, m2, d2 = map(int, m.groups())
        start = dt.date(year, m1, d1)
        end = dt.date(year + (1 if (m2, d2) < (m1, d1) else 0), m2, d2)
        out.append({'store': str(store).strip(), 'start': start, 'end': end,
                    'period': f'{m1}.{d1} ~ {m2}.{d2}'})
    return out


def load_products():
    """product-info 상품기초정보 → {(시즌, 수입품번): info}, 이미지 {품번코드: '시즌/코드.jpg'}"""
    files = sorted(glob.glob(os.path.join(PRODUCT_INFO, '상품기초정보*.xlsx')))
    if not files:
        print(f'  ! 상품기초정보 엑셀 없음: {PRODUCT_INFO}')
        return {}, {}
    images = {}
    for d in os.listdir(PRODUCT_INFO):
        p = os.path.join(PRODUCT_INFO, d)
        if os.path.isdir(p) and re.fullmatch(r'\d{2}(AW|SS)', d):
            for f in os.listdir(p):
                code, ext = os.path.splitext(f)
                if ext.lower() in ('.jpg', '.jpeg', '.png', '.webp'):
                    images[code] = f'{d}/{f}'
    wb = openpyxl.load_workbook(files[-1], read_only=True, data_only=True)
    sh = Sheet(wb.worksheets[0])
    c = {n: sh.col(n) for n in ('년도명', '시즌명', '기획명', '복종명', '수입품번', '품번', '품명', '최초가')}
    prods = {}
    for r in sh.rows:
        pn = r[c['수입품번']]
        if not pn:
            continue
        key = (season_of(r[c['년도명']], r[c['시즌명']]), str(pn).strip())
        p = prods.setdefault(key, {'name': r[c['품명']], 'plan': r[c['기획명']], 'sub': r[c['복종명']],
                                   'tag': num(r[c['최초가']]), 'codes': []})
        code = str(r[c['품번']] or '').strip()
        if code and code not in p['codes']:
            p['codes'].append(code)
    print(f'  상품기초정보: {os.path.basename(files[-1])} · 품번 {len(prods)} · 이미지 {len(images)}')
    return prods, images


def read_master(sheet, prods, images, with_io):
    """26AW / 이월 시트 → 품번별 기본정보 + SKU 목록(시트 순서)"""
    c = {n: sheet.col(n) for n in ('년도명', '시즌명', '기획명', '복종명', '수입품번', '품번', '품명', '색상', '사이즈',
                                   '최초가', '현판가', '판매구분', '최초출고일')}
    if with_io:
        c['입고'] = sheet.col('입고')
        c['판매'] = sheet.col('판매')
    items, missing = {}, set()
    for r in sheet.rows:
        pn = r[c['수입품번']]
        if not pn or not r[c['색상']]:
            continue
        pn = str(pn).strip()
        season = season_of(r[c['년도명']], r[c['시즌명']])
        if pn not in items:
            info = prods.get((season, pn))
            if info is None:
                info = next((v for (s, p), v in prods.items() if p == pn), None)
            if info is None:
                missing.add(pn)
                info = {'name': r[c['품명']], 'plan': r[c['기획명']], 'sub': r[c['복종명']],
                        'tag': num(r[c['최초가']]), 'codes': []}
            items[pn] = {'pn': pn, 'cat': cat_of(info['plan']), 'plan': info['plan'], 'sub': info['sub'],
                         'name': info['name'], 'season': season, 'tag': info['tag'] or num(r[c['최초가']]),
                         'cur': num(r[c['현판가']]), 'kind': r[c['판매구분']], 'out': date_str(r[c['최초출고일']]),
                         '_codes': list(info['codes']), '_skus': {}}
        it = items[pn]
        out = date_str(r[c['최초출고일']])
        if out and (not it['out'] or out < it['out']):
            it['out'] = out
        code = str(r[c['품번']] or '').strip()
        if code and code not in it['_codes']:
            it['_codes'].append(code)
        sk = (str(r[c['색상']]).strip(), str(r[c['사이즈']]).strip())
        io = it['_skus'].setdefault(sk, [0, 0])
        if with_io:
            io[0] += num(r[c['입고']])
            io[1] += num(r[c['판매']])
    for it in items.values():
        img = next((images[k] for k in it.pop('_codes') if k in images), None)
        it['img'] = img
    return items, missing


def build(raw):
    print(f'■ raw 파일: {os.path.basename(raw)}')
    prods, images = load_products()
    wb = openpyxl.load_workbook(raw, read_only=True, data_only=True)

    ws_cur = find_sheet(wb, lambda t: re.fullmatch(r'\d{2}(AW|SS)', t), '당시즌(예: 26AW)')
    ws_co = find_sheet(wb, lambda t: t == '이월', '이월')
    sh_sale = Sheet(find_sheet(wb, lambda t: '판매' in t and '행사' not in t, '1.판매'))
    sh_wh = Sheet(find_sheet(wb, lambda t: '창고재고' in t, '2.창고재고'))
    sh_st = Sheet(find_sheet(wb, lambda t: '매재' in t or '매장재고' in t, '3.이월매재'))
    sh_ev = Sheet(find_sheet(wb, lambda t: '행사' in t, '4.행사매장리스트'), key='매장명')
    cur_season = ws_cur.title.strip()

    # ---- 26AW / 이월 마스터
    aw_items, aw_missing = read_master(Sheet(ws_cur), prods, images, with_io=True)
    co_items, co_missing = read_master(Sheet(ws_co), prods, images, with_io=False)

    # ---- 1.판매
    c = {n: sh_sale.col(n) for n in ('판매일자', '매장형태명', '매장명', '수입품번', '색상', '사이즈', '판매수량')}
    sales = []
    for r in sh_sale.rows:
        d = to_date(r[c['판매일자']])
        if not d or not r[c['수입품번']]:
            continue
        sales.append((d, str(r[c['매장형태명']] or '').strip(), str(r[c['매장명']] or '').strip(),
                      (str(r[c['수입품번']]).strip(), str(r[c['색상']]).strip(), str(r[c['사이즈']]).strip()),
                      num(r[c['판매수량']])))
    p_start, p_end = min(s[0] for s in sales), max(s[0] for s in sales)
    print(f'  판매기간: {p_start} ~ {p_end} · 판매행 {len(sales)}')
    store_type = {}
    for d, t, st, k, q in sales:
        store_type.setdefault(st, t)

    events = parse_events(sh_ev, p_end.year)
    ev_hit = [e for e in events if e['start'] <= p_end and e['end'] >= p_start and store_type.get(e['store']) in CO_TYPES]
    ev_hit.sort(key=lambda e: (e['start'], e['store']))

    aw_sale, co_sale, co_store_sale = defaultdict(int), defaultdict(int), defaultdict(int)
    ev_qty = defaultdict(int)
    for d, t, st, k, q in sales:
        if t not in ALL_TYPES:
            continue
        aw_sale[k] += q
        if t not in CO_TYPES:
            continue
        hit = next((i for i, e in enumerate(ev_hit) if e['store'] == st and e['start'] <= d <= e['end']), None)
        if hit is not None:
            if k[0] in co_items:
                ev_qty[hit] += q
            continue
        co_sale[k] += q
        co_store_sale[(k, st)] += q

    # ---- 2.창고재고
    c = {n: sh_wh.col(n) for n in ('수입품번', '색상', '사이즈', '반영재고')}
    wh = defaultdict(int)
    for r in sh_wh.rows:
        if r[c['수입품번']]:
            wh[(str(r[c['수입품번']]).strip(), str(r[c['색상']]).strip(), str(r[c['사이즈']]).strip())] += num(r[c['반영재고']])

    def finish(items, sale_map, with_io):
        out = []
        for it in items.values():
            skus = []
            for (col, size), io in it.pop('_skus').items():
                k = (it['pn'], col, size)
                row = [col, size, sale_map.get(k, 0), wh.get(k, 0)]
                if with_io:
                    row += io
                skus.append(row)
            it['skus'] = skus
            if not with_io:
                it.pop('img', None)
            if it['cat'] and any(s[3] <= 0 for s in skus):
                out.append(it)
        return out

    aw = finish(aw_items, aw_sale, True)
    co = finish(co_items, co_sale, False)

    def rank_key(p):
        rt = [s for s in p['skus'] if s[3] <= 0]
        return (-sum(s[2] for s in rt), -sum(s[2] for s in p['skus']), p['pn'])
    aw.sort(key=rank_key)
    co.sort(key=rank_key)

    # ---- 3.이월매재 (매장 재고)
    c = {n: sh_st.col(n) for n in ('매장명', '수입품번', '색상', '사이즈')}
    ci_calc, ci_tr, ci_real = sh_st.col('수량', '전산재고'), sh_st.col('수량', '미착재고'), sh_st.col('수량', '실재고')
    st_stock = defaultdict(lambda: [0, 0, 0])  # 전산, 미착, 실
    for r in sh_st.rows:
        if not r[c['수입품번']]:
            continue
        st = str(r[c['매장명']] or '').strip()
        if store_type.get(st) not in CO_TYPES:
            continue
        k = (str(r[c['수입품번']]).strip(), str(r[c['색상']]).strip(), str(r[c['사이즈']]).strip())
        v = st_stock[(k, st)]
        v[0] += num(r[ci_calc]); v[1] += num(r[ci_tr]); v[2] += num(r[ci_real])
    stores = sorted({st for (k, st) in st_stock} | {st for (k, st) in co_store_sale})

    excl = defaultdict(list)
    for e in events:
        if e['end'] > p_end and store_type.get(e['store']) in CO_TYPES:
            excl[e['store']].append(e)
    rt_excl = [{'store': s, 'periods': [e['period'] for e in sorted(v, key=lambda e: e['start'])]}
               for s, v in sorted(excl.items())]

    rt, nomatch = [], 0
    for p in co:
        matched = []
        for col, size, sale, w in [s for s in p['skus'] if s[3] <= 0]:
            k = (p['pn'], col, size)
            rows = []
            for st in stores:
                calc, tr, real = st_stock.get((k, st), (0, 0, 0))
                rows.append({'store': st, 'sale': co_store_sale.get((k, st), 0), 'calc': calc, 'transit': tr, 'stock': real})
            inn = sorted([r for r in rows if r['sale'] > 0 and r['sale'] > r['calc']], key=lambda r: (-r['sale'], r['calc'], r['store']))[:2]
            if not inn:
                nomatch += 1
                continue
            names = {r['store'] for r in inn}
            cand = []
            for r in rows:
                if r['store'] in names or r['store'] in excl or r['sale'] >= inn[0]['sale']:
                    continue
                avail = r['stock'] - max(r['sale'], 1)
                if avail > 0:
                    cand.append((avail, r))
            if not cand:
                nomatch += 1
                continue
            avail, out = sorted(cand, key=lambda x: (-x[0], x[1]['sale'], x[1]['store']))[0]
            need = [r['sale'] - r['calc'] for r in inn]
            total = min(avail, sum(need))
            left, qtys = total, []
            for n in need:
                q = min(n, left); qtys.append(q); left -= q
            inn_out = []
            for r, q in zip(inn, qtys):
                if q <= 0:
                    continue
                e = {'store': r['store'], 'sale': r['sale'], 'stock': r['stock'], 'transit': r['transit'], 'qty': q}
                if r['sale'] == 1 and r['calc'] <= 0 and q == 1 and out['stock'] - (total + 1) >= max(out['sale'], 1) + 1:
                    e['qty'] = 2; e['plus'] = 1; total += 1
                inn_out.append(e)
            matched.append({'c': col, 's': size, 'sale': sale, 'wh': w,
                            'out': {'store': out['store'], 'sale': out['sale'], 'stock': out['stock'], 'qty': total},
                            'inn': inn_out})
        if matched:
            csum = defaultdict(int)
            for m in matched:
                csum[m['c']] += m['sale']
            size_idx = lambda s: SIZE_ORDER.index(s) if s in SIZE_ORDER else len(SIZE_ORDER)
            matched.sort(key=lambda m: (-csum[m['c']], m['c'], size_idx(m['s']), m['s']))
            rt.append({'pn': p['pn'], 'name': p['name'], 'cat': p['cat'], 'plan': p['plan'], 'sub': p['sub'], 'skus': matched})

    meta = {
        'src': os.path.basename(raw), 'season': cur_season,
        'coSeason': sorted({p['season'] for p in co}) or [],
        'start': p_start.isoformat(), 'end': p_end.isoformat(),
        'built': dt.datetime.now().strftime('%Y-%m-%d %H:%M'),
    }
    data = {'meta': meta, 'aw': aw, 'co': co,
            'ev': [{'store': e['store'], 'qty': ev_qty.get(i, 0), 'period': e['period']} for i, e in enumerate(ev_hit)],
            'rt': rt, 'rtExcl': rt_excl, 'rtInfo': {'stores': len(stores), 'nomatch': nomatch}}

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, 'w', encoding='utf-8') as f:
        f.write('window.RT_DATA=')
        json.dump(data, f, ensure_ascii=False, separators=(',', ':'))
        f.write(';\n')

    # ---- 검증 요약
    noimg = [p['pn'] for p in aw if not p['img']]
    print(f'  {cur_season} 탭: {len(aw)}개 품번 (이미지 없음 {len(noimg)}{": " + ", ".join(noimg) if noimg else ""})')
    print(f'  이월 탭: {len(co)}개 품번 · 행사 제외 {sum(ev_qty.values())}장 ({len(ev_hit)}건)')
    print(f'  이월RT매장 탭: {len(rt)}개 품번 · 매칭 SKU {sum(len(p["skus"]) for p in rt)} · 매칭 안 됨 {nomatch} · 매장 {len(stores)}')
    if aw_missing or co_missing:
        print(f'  ! 상품기초정보에 없는 품번 {len(aw_missing | co_missing)}개 (raw 시트 정보로 대체): {sorted(aw_missing | co_missing)[:15]}')
    print(f'  → {os.path.relpath(OUT, HERE)} ({os.path.getsize(OUT) // 1024} KB)')
    return data


def latest_raw():
    files = glob.glob(os.path.join(HERE, '캉골RT내역_*.xlsx')) + glob.glob(os.path.join(HERE, 'raw', '캉골RT내역_*.xlsx'))
    files = [f for f in files if not os.path.basename(f).startswith('~$')]
    if not files:
        sys.exit('캉골RT내역_YYMMDD.xlsx 파일이 없습니다.')
    return max(files, key=lambda f: (re.findall(r'(\d{6})', os.path.basename(f)) or ['0'])[-1] + str(os.path.getmtime(f)))


def make_offline(data, raw):
    """공유용 단일 HTML: 데이터와 26AW 이미지를 파일 안에 넣어 오프라인에서도 열린다."""
    import base64
    import copy
    d = copy.deepcopy(data)
    for p in d['aw']:
        path = os.path.join(PRODUCT_INFO, p['img']) if p.get('img') else None
        if path and os.path.exists(path):
            p['img'] = 'data:image/jpeg;base64,' + base64.b64encode(open(path, 'rb').read()).decode()
        else:
            p['img'] = None
    html = open(os.path.join(HERE, 'index.html'), encoding='utf-8').read()
    html = html.replace("const IMG_BASE = '../product-info/';", "const IMG_BASE = '';", 1)
    html, n = re.subn(r'\(function load\(\)\{.*?\}\)\(\);', 'boot();', html, count=1, flags=re.S)
    if n != 1:
        raise RuntimeError('index.html 로더 블록을 찾지 못함')
    payload = json.dumps(d, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
    html = html.replace('<script>\n', f'<script>window.RT_DATA={payload};</script>\n<script>\n', 1)
    tag = (re.findall(r'(\d{6})', os.path.basename(raw)) or [dt.date.today().strftime('%y%m%d')])[-1]
    out = os.path.join(HERE, 'dist', f'Kangol_RT검토맵_{tag}.html')
    os.makedirs(os.path.dirname(out), exist_ok=True)
    open(out, 'w', encoding='utf-8', newline='\n').write(html)
    print(f'  → 공유용 파일: {os.path.relpath(out, HERE)} ({os.path.getsize(out) // 1024} KB)')
    return out


def git(repo, *args):
    exe = shutil.which('git') or r'C:\Program Files\Git\cmd\git.exe'
    return subprocess.run([exe, '-C', repo, *args], capture_output=True, text=True, encoding='utf-8')


def push(repo, paths, message):
    git(repo, 'add', *paths)
    if not git(repo, 'diff', '--cached', '--quiet').returncode:
        print(f'  {os.path.basename(repo)}: 변경 없음')
        return True
    git(repo, 'commit', '-q', '-m', message)
    r = git(repo, 'push', '-q')
    print(f'  {os.path.basename(repo)}: ' + ('푸시 완료' if r.returncode == 0 else '푸시 실패\n' + r.stderr))
    return r.returncode == 0


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    args = [a for a in sys.argv[1:] if a != '--push']
    raw = args[0] if args else latest_raw()
    data = build(raw)
    if not data['aw'] or not data['co']:
        sys.exit('! 품번이 0개입니다. raw 파일 형식을 확인하세요. (푸시하지 않음)')
    make_offline(data, raw)
    if '--push' in sys.argv:
        print('■ GitHub 업로드')
        label = os.path.splitext(os.path.basename(raw))[0]
        if os.path.isdir(os.path.join(PRODUCT_INFO, '.git')):
            push(PRODUCT_INFO, ['.'], f'이미지 갱신 ({label})')
        ok = push(HERE, ['index.html', 'data/rt.js'], f'RT 데이터 업데이트 ({label})')
        if ok:
            print('  1~2분 후 반영: https://ejlim0108.github.io/RT/')
